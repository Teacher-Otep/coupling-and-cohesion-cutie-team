#!/usr/bin/env python3
"""
Generate PowerPoint presentation for Coupling and Cohesion project
Run this file to create Presentation.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_slide(prs, title, content_list):
    """Create a slide with title and bullet points"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    # Add title
    title_shape = slide.shapes.title
    title_shape.text = title
    title_frame = title_shape.text_frame
    title_frame.paragraphs[0].font.size = Pt(44)
    title_frame.paragraphs[0].font.bold = True
    
    # Add content
    body_shape = slide.placeholders[1]
    text_frame = body_shape.text_frame
    text_frame.clear()
    
    for content in content_list:
        p = text_frame.add_paragraph()
        p.text = content
        p.level = 0
        p.font.size = Pt(18)
    
    return slide

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Slide 1: Title
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
left = Inches(1)
top = Inches(2.5)
width = Inches(8)
height = Inches(2)

title_box = slide1.shapes.add_textbox(left, top, width, height)
title_frame = title_box.text_frame
title_frame.word_wrap = True

p = title_frame.paragraphs[0]
p.text = "Coffee Shop Order System Refactoring"
p.font.size = Pt(54)
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

p = title_frame.add_paragraph()
p.text = "Reducing Coupling, Increasing Cohesion"
p.font.size = Pt(32)
p.alignment = PP_ALIGN.CENTER

# Slide 2: The Problem
create_slide(prs, "The Problem", [
    "Original OrderProcessor class had 3 responsibilities:",
    "• Tax calculations",
    "• Receipt formatting",
    "• File I/O operations",
    "",
    "Result: Hard to test, modify, and extend"
])

# Slide 3: Our Solution
create_slide(prs, "Our Solution", [
    "Separated into individual classes:",
    "• Order - data model",
    "• TaxCalculator - tax logic",
    "• ReceiptFormatter - formatting",
    "• OrderRepository - persistence",
    "• OrderProcessor - orchestrator"
])

# Slide 4: Code Comparison
create_slide(prs, "Before vs After", [
    "BEFORE: 30+ lines, mixed concerns",
    "AFTER: 7 classes, single responsibilities",
    "",
    "Benefits:",
    "• Easier to test",
    "• Easier to modify",
    "• Easier to extend"
])

# Slide 5: Design Patterns
create_slide(prs, "Design Patterns Used", [
    "• Dependency Injection",
    "• Strategy Pattern (tax calculation)",
    "• Repository Pattern (storage)",
    "• Single Responsibility Principle"
])

# Slide 6: Results
create_slide(prs, "Results", [
    "Coupling: High → Low",
    "Cohesion: Low → High",
    "Testability: Poor → Good",
    "Code Quality: Improved",
    "Extensibility: Difficult → Easy"
])

# Slide 7: Team Members
create_slide(prs, "Team Members & Contributions", [
    "Gilbert Fatiig - 40%",
    "Led architecture design, Order.java, OrderProcessor.java",
    "",
    "Samo Dinnao - 30%",
    "TaxCalculator, DefaultTaxCalculator, ReceiptFormatter",
    "",
    "Eirohnjan Balino - 30%",
    "OrderRepository, FileOrderRepository, Documentation"
])

# Save presentation
prs.save('Presentation.pptx')
print("✓ Presentation.pptx created successfully!")
print("✓ Open Presentation.pptx to view the slides")
